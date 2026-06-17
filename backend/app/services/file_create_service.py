"""
File creation service — generates blank/minimal files for supported formats.
"""
import io
import hashlib
from pathlib import Path
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession
from app.models.file import Folder, File
from app.core.exceptions import NotFound, AppException
from app.config import get_settings

settings = get_settings()
UPLOAD_ROOT = Path(settings.UPLOAD_DIR).resolve()

BLANK_GENERATORS: dict[str, callable] = {}


def _register(ext: str):
    def decorator(fn):
        BLANK_GENERATORS[ext] = fn
        return fn
    return decorator


@_register("txt")
@_register("md")
@_register("json")
@_register("xml")
@_register("yaml")
@_register("yml")
@_register("csv")
@_register("log")
def _blank_text(ext: str, _mime: str) -> bytes:
    return b""


@_register("docx")
def _blank_docx(_ext: str, _mime: str) -> bytes:
    try:
        from docx import Document
        buf = io.BytesIO()
        doc = Document()
        doc.save(buf)
        return buf.getvalue()
    except ImportError:
        raise AppException("python-docx is not installed, cannot create blank docx", status_code=400)


@_register("xlsx")
def _blank_xlsx(_ext: str, _mime: str) -> bytes:
    try:
        from openpyxl import Workbook
        buf = io.BytesIO()
        wb = Workbook()
        wb.save(buf)
        return buf.getvalue()
    except ImportError:
        raise AppException("openpyxl is not installed, cannot create blank xlsx", status_code=400)


@_register("pptx")
def _blank_pptx(_ext: str, _mime: str) -> bytes:
    try:
        from pptx import Presentation
        buf = io.BytesIO()
        prs = Presentation()
        prs.save(buf)
        return buf.getvalue()
    except ImportError:
        raise AppException("python-pptx is not installed, cannot create blank pptx", status_code=400)


def _get_mime(ext: str) -> str:
    mime_map = {
        "txt": "text/plain", "md": "text/markdown", "json": "application/json",
        "xml": "application/xml", "yaml": "text/yaml", "yml": "text/yaml",
        "csv": "text/csv", "log": "text/plain",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }
    return mime_map.get(ext, "application/octet-stream")


async def create_file(
    db: AsyncSession,
    name: str,
    extension: str,
    owner_id: int,
    folder_id: int | None = None,
) -> dict:
    """Create a blank file with the given extension."""
    ext = extension.lower().lstrip(".")
    if ext not in BLANK_GENERATORS:
        raise AppException(f"Unsupported file format for creation: .{ext}", status_code=400)

    # Validate target folder
    target_folder_id = folder_id if (folder_id and folder_id > 0) else None
    if target_folder_id:
        folder = await db.get(Folder, target_folder_id)
        if not folder or folder.deleted:
            raise NotFound("Target folder not found")
        if folder.owner_id != owner_id:
            raise AppException("Access denied: target folder does not belong to current user", status_code=403)

    # Auto-name to avoid conflicts
    base_name = name
    final_name = name
    counter = 1
    while True:
        existing = await db.execute(
            select(File).where(
                File.name == final_name,
                File.extension == ext,
                File.folder_id == target_folder_id,
                File.owner_id == owner_id,
                File.deleted == False,
            )
        )
        if not existing.scalar_one_or_none():
            break
        counter += 1
        final_name = f"{base_name} {counter}"

    mime_type = _get_mime(ext)
    content = BLANK_GENERATORS[ext](ext, mime_type)
    md5_hash = hashlib.md5(content).hexdigest()
    content_path = f"{md5_hash[:2]}/{md5_hash[2:4]}/{md5_hash}.{ext}" if ext else f"{md5_hash[:2]}/{md5_hash[2:4]}/{md5_hash}"
    abs_content_path = UPLOAD_ROOT / content_path

    # Check for existing content (dedup)
    existing_content = await db.execute(
        select(File).where(File.md5_hash == md5_hash, File.deleted == False).limit(1)
    )
    existing_file = existing_content.scalar_one_or_none()

    deduplicated = False
    if existing_file and abs_content_path.exists():
        deduplicated = True
        new_file = File(
            name=final_name, extension=ext, size=len(content),
            folder_id=target_folder_id, owner_id=owner_id,
            storage_path=existing_file.storage_path,
            mime_type=mime_type, md5_hash=md5_hash, ref_count=1, deleted=False,
        )
        db.add(new_file)
        await db.flush()
        await db.execute(update(File).where(File.id == existing_file.id).values(ref_count=File.ref_count + 1))
        await db.commit()
        await db.refresh(new_file)
    else:
        abs_content_path.parent.mkdir(parents=True, exist_ok=True)
        abs_content_path.write_bytes(content)
        new_file = File(
            name=final_name, extension=ext, size=len(content),
            folder_id=target_folder_id, owner_id=owner_id,
            storage_path=content_path,
            mime_type=mime_type, md5_hash=md5_hash, ref_count=1, deleted=False,
        )
        db.add(new_file)
        await db.commit()
        await db.refresh(new_file)

    return {
        "id": new_file.id, "name": new_file.name, "extension": new_file.extension,
        "size": new_file.size, "mime_type": new_file.mime_type, "deduplicated": deduplicated,
    }
